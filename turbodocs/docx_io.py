"""Convert between .docx files and QTextDocument tabs.

Tabs are stored in the docx as bookmarks (`tdtab_<n>`) with titles kept in the core-properties comment,
so files stay plain, valid Word documents that any editor can open.
"""
import io
import json
import zipfile

from docx import Document
from docx.document import Document as _DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Emu, Pt, RGBColor
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt
from PySide6.QtGui import (QColor, QFont, QFontDatabase, QImage, QTextCharFormat, QTextCursor, QTextFormat,
                           QTextLength, QTextListFormat, QTextTableFormat)

from . import images, styles, theme
from .document import new_document
from .styles import DocStyles, StyleDef

EMU_PX = 9525
MAX_EXPORT_IMG_PX = 624  # 6.5in at 96 dpi
TEXT_WIDTH_EMU = 5943600
TAB_MARK = "tdtab_"
TAB_META = "turbodocs-tabs:"
HIGHLIGHT_FG = "#161616"
CONTRAST_LIGHT = "#f2f2f2"   # display-only text colours placed on shaded cells / highlights
DISPLAY_ONLY = (HIGHLIGHT_FG, CONTRAST_LIGHT)


def contrast_for(color) -> str:
    """Readable display-only text colour for a background (dark on light, light on dark)."""
    return HIGHLIGHT_FG if QColor(color).lightnessF() > 0.55 else CONTRAST_LIGHT
ALIGN_TO_QT = {
    WD_ALIGN_PARAGRAPH.CENTER: Qt.AlignmentFlag.AlignHCenter,
    WD_ALIGN_PARAGRAPH.RIGHT: Qt.AlignmentFlag.AlignRight,
    WD_ALIGN_PARAGRAPH.JUSTIFY: Qt.AlignmentFlag.AlignJustify,
}
HL_NAMED = {
    WD_COLOR_INDEX.YELLOW: "#ffff00", WD_COLOR_INDEX.BRIGHT_GREEN: "#00ff00", WD_COLOR_INDEX.TURQUOISE: "#00ffff",
    WD_COLOR_INDEX.PINK: "#ff00ff", WD_COLOR_INDEX.BLUE: "#0000ff", WD_COLOR_INDEX.RED: "#ff0000",
    WD_COLOR_INDEX.DARK_BLUE: "#000080", WD_COLOR_INDEX.TEAL: "#008080", WD_COLOR_INDEX.GREEN: "#008000",
    WD_COLOR_INDEX.VIOLET: "#800080", WD_COLOR_INDEX.DARK_RED: "#800000", WD_COLOR_INDEX.DARK_YELLOW: "#808000",
    WD_COLOR_INDEX.GRAY_50: "#808080", WD_COLOR_INDEX.GRAY_25: "#c0c0c0",
}


# ---------------------------------------------------------------- shared table helpers

def equal_pcts(n):
    return [100.0 / n] * n


def table_format(ncols, pcts=None) -> QTextTableFormat:
    pcts = pcts if pcts and len(pcts) == ncols else equal_pcts(ncols)
    total = sum(pcts) or 100.0
    tf = QTextTableFormat()
    tf.setBorder(1)
    tf.setBorderBrush(QColor(theme.mix(theme.cur["bg"], theme.cur["text"], 0.3)))
    tf.setBorderStyle(QTextTableFormat.BorderStyle.BorderStyle_Solid)
    tf.setBorderCollapse(True)
    tf.setCellPadding(7)
    tf.setCellSpacing(0)
    tf.setWidth(QTextLength(QTextLength.Type.PercentageLength, 100))
    tf.setColumnWidthConstraints([QTextLength(QTextLength.Type.PercentageLength, p * 100.0 / total) for p in pcts])
    tf.setTopMargin(6)
    tf.setBottomMargin(10)
    return tf


def shade_cell(table, cell, color: str):
    """Set (or clear, with '') a table cell's background and keep its text readable on it."""
    fmt = cell.format().toTableCellFormat()
    if color:
        fmt.setBackground(QColor(color))
    else:
        fmt.clearBackground()
    cell.setFormat(fmt)
    want = None
    if color:
        want = HIGHLIGHT_FG if QColor(color).lightnessF() > 0.55 else CONTRAST_LIGHT
    doc = table.document()
    start, end = cell.firstPosition(), cell.lastPosition()
    block = doc.findBlock(start)
    c = QTextCursor(doc)
    while block.isValid() and block.position() <= end:
        spans = [(f.position(), f.position() + f.length(), QTextCharFormat(f.charFormat()))
                 for f in (it.fragment() for it in block)]
        bf = QTextCursor(block)
        for a, b, f in spans:
            if f.hasProperty(QTextFormat.Property.ForegroundBrush) and not f.isAnchor() and \
                    f.foreground().color().name() not in DISPLAY_ONLY:
                continue  # an explicit colour chosen by the user stays
            if f.isAnchor():
                continue
            f.clearForeground()
            if want:
                f.setForeground(QColor(want))
            c.setPosition(a)
            c.setPosition(b, QTextCursor.MoveMode.KeepAnchor)
            c.setCharFormat(f)
        bcf = bf.blockCharFormat()
        bcf.clearForeground()
        if want:
            bcf.setForeground(QColor(want))
        bf.setBlockCharFormat(bcf)
        block = block.next()


def col_pcts(table) -> list:
    cons = table.format().columnWidthConstraints()
    n = table.columns()
    if len(cons) == n and all(c.type() == c.Type.PercentageLength for c in cons):
        vals = [c.rawValue() for c in cons]
        total = sum(vals) or 100.0
        return [v * 100.0 / total for v in vals]
    return equal_pcts(n)


# ---------------------------------------------------------------- import

_FROM_DOCX_LOWER = {k.lower(): v for k, v in styles.FROM_DOCX.items()}

def _iter_blocks(parent):
    elm = parent.element.body if isinstance(parent, _DocxDocument) else parent._tc
    for child in elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def _chain(style, getter):
    while style is not None:
        v = getter(style)
        if v is not None:
            return v
        style = style.base_style
    return None


def _find_style(dx, name):
    for s in dx.styles:
        if s.name and s.name.lower() == name.lower():
            return s
    return None


def read_styles(dx) -> DocStyles:
    st = DocStyles()
    normal = _find_style(dx, "Normal")
    if normal is None:
        return st
    bg = dx.element.find(qn("w:background"))
    if bg is not None:
        col = bg.get(qn("w:color")) or ""
        if len(col) == 6 and col.lower() not in ("ffffff", "auto"):
            st.page_color = "#" + col.lower()
    fam = _chain(normal, lambda s: s.font.name)
    size = _chain(normal, lambda s: s.font.size)
    if fam:
        st.base_family = fam
    if size:
        st.base_size = size.pt
    for key in styles.STYLED:
        s = _find_style(dx, styles.DOCX_NAME[key])
        if s is None:
            continue
        d = st.defs[key]
        if key != "Normal":
            sz = _chain(s, lambda x: x.font.size)
            if sz:
                d.size = sz.pt
            fam_k = _chain(s, lambda x: x.font.name)
            if fam_k and fam_k != st.base_family:
                d.family = fam_k
        b = _chain(s, lambda x: x.font.bold)
        i = _chain(s, lambda x: x.font.italic)
        if key != "Normal":
            d.bold, d.italic = bool(b), bool(i)
        try:
            rgb = _chain(s, lambda x: x.font.color.rgb if x.font.color is not None else None)
        except Exception:
            rgb = None
        if rgb is not None and theme.contrast_ok(QColor("#" + str(rgb))):
            d.color = "#" + str(rgb)
        pf = s.paragraph_format
        bef = _chain(s, lambda x: x.paragraph_format.space_before)
        aft = _chain(s, lambda x: x.paragraph_format.space_after)
        line = _chain(s, lambda x: x.paragraph_format.line_spacing)
        if bef is not None:
            d.before = bef / EMU_PX
        if aft is not None:
            d.after = aft / EMU_PX
        if isinstance(line, float):
            d.line = line
    return st


def _load_embedded_fonts(path):
    try:
        with zipfile.ZipFile(path) as z:
            for n in z.namelist():
                if n.lower().startswith("word/fonts/") and n.lower().endswith((".ttf", ".otf")):
                    QFontDatabase.addApplicationFontFromData(QByteArray(z.read(n)))
    except (zipfile.BadZipFile, OSError):
        pass


def _list_kind(dx, p, sname):
    """-> (kind, level) or (None, 0)"""
    for prefix, kind in (("List Bullet", "bullet"), ("List Number", "number")):
        if sname.startswith(prefix):
            tail = sname[len(prefix):].strip()
            return kind, int(tail) if tail.isdigit() else 1
    pPr = p._p.pPr
    numPr = pPr.numPr if pPr is not None else None
    if numPr is None or numPr.numId is None or numPr.numId.val == 0:
        return None, 0
    ilvl = numPr.ilvl.val if numPr.ilvl is not None else 0
    try:
        numbering = dx.part.numbering_part.element
        abstract = numbering.xpath(f'w:num[@w:numId="{numPr.numId.val}"]/w:abstractNumId/@w:val')
        fmt = numbering.xpath(
            f'w:abstractNum[@w:abstractNumId="{abstract[0]}"]/w:lvl[@w:ilvl="{ilvl}"]/w:numFmt/@w:val')
        return ("bullet" if not fmt or fmt[0] == "bullet" else "number"), ilvl + 1
    except Exception:
        return "bullet", ilvl + 1


def _is_break(item):
    return isinstance(item, Paragraph) and item._p.pPr is not None and item._p.pPr.find(qn("w:sectPr")) is not None


def _split_sections(items):
    """Google Docs exports each tab as a Word section: split where a paragraph carries a section break.
    A paragraph that only carries the break is dropped so tabs don't gain a stray blank line."""
    out, current = [], []
    for item in items:
        if _is_break(item):
            if item.text.strip() or not current:
                current.append(item)          # real text, or an otherwise empty tab keeping one empty paragraph
            out.append(current)
            current = []
        else:
            current.append(item)
    if current and (not out or any(not isinstance(i, Paragraph) or i.text.strip() for i in current)):
        out.append(current)
    elif current:
        out[-1].extend(current)
    return out or [items]


def _tab_title(items):
    """Google's .docx export carries no tab names, so guess: a Title, else a heading, else a short first line."""
    paras = [i for i in items if isinstance(i, Paragraph) and i.text.strip()]

    def style(p):
        try:
            return (p.style.name or "").lower()
        except Exception:
            return ""
    for want in (("title",), ("heading 1", "heading 2", "heading 3")):
        for p in paras:
            if style(p) in want:
                return p.text.strip()[:40]
    for p in paras[:1]:
        t = p.text.strip()
        if len(t) <= 32:
            return t
    return ""


def _split_tabs(dx):
    items, tabs, current = list(_iter_blocks(dx)), [], []
    for item in items:
        if isinstance(item, Paragraph) and item._p.xpath(f'.//w:bookmarkStart[starts-with(@w:name,"{TAB_MARK}")]'):
            if current:
                tabs.append(current)
            current = []
        current.append(item)
    tabs.append(current)
    if len(tabs) == 1:
        tabs = _split_sections(tabs[0])
    titles = []
    comments = dx.core_properties.comments or ""
    if comments.startswith(TAB_META):
        try:
            titles = [str(t) for t in json.loads(comments[len(TAB_META):])]
        except ValueError:
            pass
    return [(titles[i] if i < len(titles) and titles[i] else (_tab_title(t) if len(tabs) > 1 else "") or f"Tab {i + 1}", t)
            for i, t in enumerate(tabs)]


class _Importer:
    def __init__(self, qdoc, dx, st):
        self.qdoc, self.dx, self.st = qdoc, dx, st
        self.lists = {}

    def run(self, items):
        c = QTextCursor(self.qdoc)
        fresh = True
        for item in items:
            if isinstance(item, Paragraph):
                if not fresh:
                    c.insertBlock()
                self.paragraph(c, item)
                fresh = False
            else:
                self.lists.clear()
                self.table(c, item)
                c.setPosition(self.qdoc.characterCount() - 1)  # the trailing block after the table
                fresh = True
        if fresh and not items:
            c.setBlockFormat(self.st.block_format("Normal"))

    def paragraph(self, c, p):
        sname = p.style.name if p.style is not None else "Normal"
        key = _FROM_DOCX_LOWER.get(sname.lower(), "Normal")
        kind, level = _list_kind(self.dx, p, sname)
        if kind:
            key = "Bulleted list" if kind == "bullet" else "Numbered list"
        bf = self.st.block_format(key)
        if p.alignment in ALIGN_TO_QT:
            bf.setAlignment(ALIGN_TO_QT[p.alignment])
        pf = p.paragraph_format
        if pf.left_indent and not kind:
            bf.setLeftMargin(pf.left_indent / EMU_PX)
        if pf.first_line_indent and not kind:
            bf.setTextIndent(pf.first_line_indent / EMU_PX)
        if pf.space_before is not None:
            bf.setTopMargin(pf.space_before / EMU_PX)
        if pf.space_after is not None:
            bf.setBottomMargin(pf.space_after / EMU_PX)
        if isinstance(pf.line_spacing, float):
            bf.setLineHeight(pf.line_spacing * 100, bf.LineHeightTypes.ProportionalHeight.value)
        c.setBlockFormat(bf)
        base = self.st.char_format(key)
        c.setBlockCharFormat(base)

        if kind:
            for lv in [k for k in self.lists if k > level]:
                del self.lists[lv]
            cur = self.lists.get(level)
            if cur and cur[0] == kind:
                cur[1].add(c.block())
            else:
                self.lists[level] = (kind, c.createList(styles.list_format(key, level)))
        else:
            self.lists.clear()

        inner = p.iter_inner_content() if hasattr(p, "iter_inner_content") else p.runs
        for item in inner:
            if isinstance(item, Run):
                self.run_content(c, item, base, None)
            else:
                for r in item.runs:
                    self.run_content(c, r, base, item.url)

    def run_content(self, c, r, base, url):
        fmt = QTextCharFormat(base)
        f = r.font
        if r.bold is not None:
            fmt.setFontWeight(QFont.Weight.Bold if r.bold else QFont.Weight.Normal)
        if r.italic is not None:
            fmt.setFontItalic(bool(r.italic))
        if r.underline:
            fmt.setFontUnderline(True)
        if f.strike:
            fmt.setFontStrikeOut(True)
        if f.name:
            fmt.setFontFamilies([f.name])
        if f.size:
            fmt.setFontPointSize(f.size.pt)
        hl = None
        try:
            if f.highlight_color in HL_NAMED:
                hl = HL_NAMED[f.highlight_color]
        except Exception:
            pass
        fill = r._r.xpath("w:rPr/w:shd/@w:fill")
        if hl is None and fill and fill[0] not in ("auto", "FFFFFF", "ffffff"):
            hl = "#" + fill[0]
        try:
            rgb = f.color.rgb if f.color is not None and f.color.type is not None else None
        except Exception:
            rgb = None
        if hl:
            fmt.setBackground(QColor(hl))
        if rgb is not None and (hl or theme.contrast_ok(QColor("#" + str(rgb)))):
            fmt.setForeground(QColor("#" + str(rgb)))
        elif hl:
            fmt.setForeground(QColor(contrast_for(hl)))
        if url:
            fmt.setAnchor(True)
            fmt.setAnchorHref(url)
            fmt.setFontUnderline(True)
            fmt.setForeground(QColor(theme.cur["accent"]))

        text = []

        def flush():
            if text:
                c.insertText("".join(text), fmt)
                text.clear()

        for child in r._r.iterchildren():
            tag = child.tag
            if tag == qn("w:t"):
                text.append(child.text or "")
            elif tag == qn("w:tab"):
                text.append("\t")
            elif tag in (qn("w:br"), qn("w:cr")):
                if child.get(qn("w:type")) != "page":
                    text.append(" ")
            elif tag == qn("w:noBreakHyphen"):
                text.append("‑")
            elif tag == qn("w:drawing"):
                flush()
                self.image(c, child)
        flush()

    def image(self, c, drawing):
        blips = drawing.xpath(".//a:blip")
        if not blips:
            return
        try:
            blob = self.dx.part.related_parts[blips[0].get(qn("r:embed"))].blob
        except Exception:
            return
        img = QImage.fromData(blob)
        if img.isNull():
            return
        ext = drawing.xpath(".//wp:extent")
        w = int(ext[0].get("cx")) / EMU_PX if ext else img.width()
        h = int(ext[0].get("cy")) / EMU_PX if ext else img.height()
        crop = [0.0, 0.0, 0.0, 0.0]
        for i, a in enumerate("ltrb"):
            v = drawing.xpath(f".//pic:blipFill/a:srcRect/@{a}")
            if v:
                crop[i] = max(0.0, int(v[0]) / 100000.0)
        key = images.register(img, blob, "jpg" if blob[:2] == b"\xff\xd8" else "png")
        c.insertText("￼", images.make_format(key, w, h, tuple(crop)))

    def table(self, c, t):
        rows = t.rows
        grid = [int(float(g.get(qn("w:w"), 0) or 0)) for g in t._tbl.xpath("w:tblGrid/w:gridCol")]
        ncols = len(grid)
        if not rows or not ncols:
            return
        pcts = grid if all(grid) else None
        tf = table_format(ncols, pcts)
        tblW = t._tbl.tblPr.find(qn("w:tblW")) if t._tbl.tblPr is not None else None
        if tblW is not None and tblW.get(qn("w:type")) == "pct":
            try:
                tf.setWidth(QTextLength(QTextLength.Type.PercentageLength, max(25.0, min(100.0, float(tblW.get(qn("w:w"))) / 50.0))))
            except (TypeError, ValueError):
                pass
        qt = c.insertTable(len(rows), ncols, tf)
        spans, shaded = [], []
        for ri, row in enumerate(rows):
            ci = 0
            for tc in row._tr.tc_lst:
                span = tc.grid_span
                if ci >= ncols:
                    break
                self.cell(qt.cellAt(ri, ci).firstCursorPosition(), _Cell(tc, t))
                fill = tc.xpath("w:tcPr/w:shd/@w:fill")
                if fill and fill[0].lower() not in ("auto", "ffffff"):
                    shaded.append((ri, ci, "#" + fill[0].lower()))
                if span > 1:
                    spans.append((ri, ci, min(span, ncols - ci)))
                ci += span
        for ri, ci, span in spans:
            qt.mergeCells(ri, ci, 1, span)
        for ri, ci, color in shaded:
            shade_cell(qt, qt.cellAt(ri, ci), color)

    def cell(self, cc, cell):
        saved, self.lists = self.lists, {}
        fresh = True
        for item in _iter_blocks(cell):
            if not isinstance(item, Paragraph):
                continue
            if not fresh:
                cc.insertBlock()
            self.paragraph(cc, item)
            fresh = False
        self.lists = saved


def import_docx(path):
    """-> (DocStyles, [(title, QTextDocument), ...])"""
    dx = Document(path)
    _load_embedded_fonts(path)
    st = read_styles(dx)
    out = []
    for title, items in _split_tabs(dx):
        doc = new_document(st, title)
        doc.setUndoRedoEnabled(False)
        _Importer(doc, dx, st).run(items)
        doc.setUndoRedoEnabled(True)
        doc.setModified(False)
        out.append((title, doc))
    return st, out


# ---------------------------------------------------------------- export

def _px_emu(px):
    return Emu(int(px * EMU_PX))


def _set_font_name(rpr_owner_font, name):
    rpr_owner_font.name = name
    rpr = rpr_owner_font._element.rPr
    if rpr is None:
        return
    rf = rpr.find(qn("w:rFonts"))
    if rf is not None:
        for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(a)) is not None:
                del rf.attrib[qn(a)]
        for a in ("w:hAnsi", "w:cs", "w:eastAsia"):
            rf.set(qn(a), name)


def _add_shading(run, hex_fill):
    rPr = run._r.get_or_add_rPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill.lstrip("#").upper())
    rPr.insert_element_before(shd, "w:fitText", "w:vertAlign", "w:rtl", "w:cs", "w:em", "w:lang",
                              "w:eastAsianLayout", "w:specVanish", "w:oMath")


def _add_hyperlink(par, url, text, fmt):
    rid = par.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    par._p.append(h)
    run = Run(OxmlElement("w:r"), par)
    h.append(run._r)
    run.text = text
    run.font.underline = True
    run.font.color.rgb = RGBColor(0x0B, 0x6B, 0xCB)
    _apply_fonts(run, fmt, None, link=True)


def _apply_fonts(run, fmt, style_def, link=False):
    if fmt.fontWeight() >= QFont.Weight.DemiBold:
        run.bold = True
    if fmt.fontItalic():
        run.italic = True
    if fmt.fontUnderline() and not fmt.isAnchor():
        run.underline = True
    if fmt.fontStrikeOut():
        run.font.strike = True
    if fmt.hasProperty(QTextFormat.Property.FontFamilies):
        fams = fmt.property(QTextFormat.Property.FontFamilies)
        if fams:
            _set_font_name(run.font, fams[0])
    if fmt.hasProperty(QTextFormat.Property.FontPointSize) and fmt.fontPointSize() > 0:
        run.font.size = Pt(fmt.fontPointSize())
    highlighted = fmt.hasProperty(QTextFormat.Property.BackgroundBrush) and fmt.background().style() != Qt.BrushStyle.NoBrush
    if highlighted and not link:
        _add_shading(run, fmt.background().color().name())
    if not link and fmt.hasProperty(QTextFormat.Property.ForegroundBrush) and not fmt.isAnchor():
        col = fmt.foreground().color()
        if col.name() in DISPLAY_ONLY:
            return  # display-only contrast colour
        if style_def is not None and style_def.color and col.name() == QColor(style_def.color).name():
            return
        if col.name() != theme.cur["text"]:
            run.font.color.rgb = RGBColor(col.red(), col.green(), col.blue())


class _Exporter:
    def __init__(self, tabs, st: DocStyles):
        self.tabs, self.st = tabs, st
        self.dx = Document()
        self.qdoc = None
        sec = self.dx.sections[0]
        sec.left_margin = sec.right_margin = Emu(914400)
        self.write_styles()
        if st.page_color:
            bg = OxmlElement("w:background")
            bg.set(qn("w:color"), st.page_color.lstrip("#").upper())
            self.dx.element.insert(0, bg)
            settings = self.dx.settings.element
            flag = OxmlElement("w:displayBackgroundShape")
            early = ("w:writeProtection", "w:view", "w:zoom", "w:removePersonalInformation",
                     "w:removeDateAndTime", "w:doNotDisplayPageBoundaries")
            anchor = [settings.find(qn(t)) for t in early if settings.find(qn(t)) is not None]
            anchor[-1].addnext(flag) if anchor else settings.insert(0, flag)

    def write_styles(self):
        dx, st = self.dx, self.st
        for key in styles.STYLED:
            try:
                s = dx.styles[styles.DOCX_NAME[key]]
            except KeyError:
                continue
            d = st.defs[key]
            if key == "Normal":
                _set_font_name(s.font, st.base_family)
                s.font.size = Pt(st.base_size)
            else:
                s.font.size = Pt(d.size or st.base_size)
                s.font.bold = d.bold
                s.font.italic = d.italic
                _set_font_name(s.font, d.family or st.base_family)
            s.font.color.rgb = RGBColor.from_string(d.color.lstrip("#").upper()) if d.color else None
            s.paragraph_format.space_before = Pt(d.before / 1.3333)
            s.paragraph_format.space_after = Pt(d.after / 1.3333)
            s.paragraph_format.line_spacing = d.line

    def run(self):
        titles = []
        for i, (title, qdoc) in enumerate(self.tabs):
            self.qdoc = qdoc
            titles.append(title)
            self.first_of_tab = i
            self.pending_mark = i
            self.blocks(qdoc.begin(), None, None, self.dx)
            if self.pending_mark is not None:  # empty tab
                self.mark(self.dx.add_paragraph())
        if len(self.tabs) > 1:
            self.dx.core_properties.comments = TAB_META + json.dumps(titles)
        if not self.dx.paragraphs and not self.dx.tables:
            self.dx.add_paragraph()
        return self.dx

    def mark(self, par):
        i = self.pending_mark
        self.pending_mark = None
        s, e = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
        s.set(qn("w:id"), str(100 + i))
        s.set(qn("w:name"), f"{TAB_MARK}{i}")
        e.set(qn("w:id"), str(100 + i))
        pPr = par._p.pPr
        par._p.insert(0 if pPr is None else 1, s)
        par._p.insert(1 if pPr is None else 2, e)

    def blocks(self, block, end_pos, cell, container):
        first = True
        while block.isValid() and (end_pos is None or block.position() <= end_pos):
            table = QTextCursor(block).currentTable() if cell is None else None
            if table is not None:
                if self.pending_mark is not None:
                    self.mark(container.add_paragraph())
                self.table(table, container)
                block = self.qdoc.findBlock(table.lastPosition()).next()
                first = False
                continue
            par = container.paragraphs[0] if (cell is not None and first) else container.add_paragraph()
            first = False
            self.paragraph(block, par)
            if cell is None and self.pending_mark is not None:
                self.mark(par)
            block = block.next()

    def paragraph(self, block, par):
        bf = block.blockFormat()
        key = styles.key_of_block(block)
        lst = block.textList()
        sname = styles.DOCX_NAME[key]
        if lst is not None:
            level = min(max(lst.format().indent(), 1), 3)
            sname += f" {level}" if level > 1 else ""
        par.style = self.dx.styles[sname]
        a = bf.alignment()
        if a & Qt.AlignmentFlag.AlignHCenter:
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif a & Qt.AlignmentFlag.AlignRight:
            par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        elif a & Qt.AlignmentFlag.AlignJustify:
            par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = par.paragraph_format
        left = bf.leftMargin() + bf.indent() * self.qdoc.indentWidth()
        if left > 1 and lst is None:
            pf.left_indent = _px_emu(left)
        if bf.textIndent() and lst is None:
            pf.first_line_indent = _px_emu(bf.textIndent())
        pf.space_before = Pt(bf.topMargin() * 0.75)
        pf.space_after = Pt(bf.bottomMargin() * 0.75)
        if bf.lineHeightType() == bf.LineHeightTypes.ProportionalHeight.value and bf.lineHeight() > 0:
            pf.line_spacing = bf.lineHeight() / 100.0
        sdef = self.st.def_for(key)

        for it in block:
            frag = it.fragment()
            fmt = frag.charFormat()
            if images.is_image(fmt):
                self.image(par, fmt)
                continue
            text = frag.text().replace("￼", "").replace(" ", "\n")
            if not text:
                continue
            if fmt.isAnchor() and fmt.anchorHref():
                _add_hyperlink(par, fmt.anchorHref(), text, fmt)
            else:
                run = par.add_run(text)
                _apply_fonts(run, fmt, sdef)

    def image(self, par, fmt):
        key, w, h, crop = images.read_format(fmt)
        entry = images.REGISTRY.get(key)
        if entry is None:
            return
        blob = entry.blob
        if blob is None:
            ba = QByteArray()
            buf = QBuffer(ba)
            buf.open(QIODevice.OpenModeFlag.WriteOnly)
            entry.img.save(buf, "PNG")
            blob = entry.blob = ba.data()
            entry.ext = "png"
        if w > MAX_EXPORT_IMG_PX:
            h, w = h * MAX_EXPORT_IMG_PX / w, MAX_EXPORT_IMG_PX
        run = par.add_run()
        run.add_picture(io.BytesIO(blob), width=_px_emu(w), height=_px_emu(h))
        if any(crop):
            blip = run._r.xpath(".//pic:blipFill/a:blip")[0]
            src = OxmlElement("a:srcRect")
            for name, v in zip("ltrb", crop):
                if v > 0:
                    src.set(name, str(int(v * 100000)))
            blip.addnext(src)

    def table(self, qt, container):
        rows, cols = qt.rows(), qt.columns()
        dt = container.add_table(rows=rows, cols=cols)
        dt.style = self.dx.styles["Table Grid"]
        dt.autofit = False
        pcts = col_pcts(qt)
        widths = [_px_emu(0) for _ in pcts]
        total = TEXT_WIDTH_EMU if container is self.dx else TEXT_WIDTH_EMU // 2
        widths = [Emu(int(total * p / 100.0)) for p in pcts]
        for i, wdt in enumerate(widths):
            dt.columns[i].width = wdt
        tblW = dt._tbl.tblPr.find(qn("w:tblW"))
        if tblW is not None:
            wl = qt.format().width()
            pct = wl.rawValue() if wl.type() == wl.Type.PercentageLength else 100.0
            tblW.set(qn("w:type"), "pct")
            tblW.set(qn("w:w"), str(int(pct * 50)))
        for r in range(rows):
            for c in range(cols):
                cell = qt.cellAt(r, c)
                if cell.row() != r or cell.column() != c:
                    continue
                dcell = dt.cell(r, c)
                rs, cs = cell.rowSpan(), cell.columnSpan()
                if rs > 1 or cs > 1:
                    dcell = dcell.merge(dt.cell(r + rs - 1, c + cs - 1))
                dcell.width = Emu(sum(widths[c:c + cs]))
                bgb = cell.format().background()
                if bgb.style() != Qt.BrushStyle.NoBrush:
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), bgb.color().name().lstrip("#").upper())
                    dcell._tc.get_or_add_tcPr().insert_element_before(
                        shd, "w:noWrap", "w:tcMar", "w:textDirection", "w:tcFitText", "w:vAlign", "w:hideMark")
                start, end = cell.firstPosition(), cell.lastPosition()
                self.blocks(self.qdoc.findBlock(start), end, dcell, dcell)


def export_docx(tabs, st: DocStyles, path):
    """tabs: [(title, QTextDocument)]"""
    _Exporter(tabs, st).run().save(path)
